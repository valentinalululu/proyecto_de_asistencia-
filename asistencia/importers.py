"""Lectura y validación de los Excel (RF-11, RF-12). Devuelve (datos, errores)."""
import re
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _leer_filas(archivo):
    """Devuelve [(nro_fila_excel, valores)] sin la cabecera ni filas vacías."""
    wb = load_workbook(archivo, data_only=True, read_only=True)
    ws = wb.active
    filas = []
    for i, valores in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue  # cabecera
        if all(v is None or str(v).strip() == "" for v in valores):
            continue
        filas.append((i, valores))
    return filas


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def parse_estudiantes(archivo):
    try:
        filas = _leer_filas(archivo)
    except Exception:
        return [], ["No se pudo leer el archivo de estudiantes. Verifica que sea un .xlsx válido."]

    errores, datos = [], []
    registros, correos = {}, {}
    for n, v in filas:
        v = list(v) + [None] * 3
        nro, nombre, correo = _texto(v[0]), _texto(v[1]), _texto(v[2]).lower()
        problemas = []
        if not nro:
            problemas.append("falta el número de registro")
        if not nombre:
            problemas.append("falta el nombre completo")
        if not correo:
            problemas.append("falta el correo electrónico")
        elif not EMAIL_RE.match(correo):
            problemas.append(f"correo inválido ({correo})")
        if nro and nro in registros:
            problemas.append(f"número de registro {nro} duplicado (ya aparece en la fila {registros[nro]})")
        if correo and EMAIL_RE.match(correo) and correo in correos:
            problemas.append(f"correo {correo} duplicado (ya aparece en la fila {correos[correo]})")
        if problemas:
            errores.append(f"Fila {n}: " + "; ".join(problemas))
            continue
        registros[nro], correos[correo] = n, n
        datos.append({"nro_registro": nro, "nombre": nombre, "correo": correo})

    if not filas:
        errores.append("El archivo de estudiantes no contiene filas de datos.")
    return datos, errores


def _a_fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            pass
    return None


def _a_hora(valor):
    if isinstance(valor, datetime):
        return valor.time().replace(second=0, microsecond=0)
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    if isinstance(valor, (int, float)) and 0 <= valor < 1:  # hora como fracción de día de Excel
        minutos = round(valor * 24 * 60)
        return (datetime.min + timedelta(minutes=minutos)).time()
    texto = _texto(valor)
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(texto, fmt).time().replace(second=0)
        except ValueError:
            pass
    return None


def parse_fechas(archivo):
    try:
        filas = _leer_filas(archivo)
    except Exception:
        return [], ["No se pudo leer el archivo de fechas. Verifica que sea un .xlsx válido."]

    errores, datos, vistas = [], [], {}
    for n, v in filas:
        v = list(v) + [None] * 3
        fecha, inicio, fin = _a_fecha(v[0]), _a_hora(v[1]), _a_hora(v[2])
        problemas = []
        if fecha is None:
            problemas.append("fecha inválida (usa dd/mm/aaaa)")
        if inicio is None:
            problemas.append("hora de inicio inválida (usa HH:MM)")
        if fin is None:
            problemas.append("hora de fin inválida (usa HH:MM)")
        if inicio and fin and fin <= inicio:
            problemas.append("la hora de fin debe ser posterior a la de inicio")
        if fecha and fecha in vistas:
            problemas.append(f"fecha {fecha:%d/%m/%Y} duplicada (ya aparece en la fila {vistas[fecha]})")
        if problemas:
            errores.append(f"Fila {n}: " + "; ".join(problemas))
            continue
        vistas[fecha] = n
        datos.append({"fecha": fecha.isoformat(), "inicio": inicio.strftime("%H:%M"), "fin": fin.strftime("%H:%M")})

    if not filas:
        errores.append("El archivo de fechas no contiene filas de datos.")
    datos.sort(key=lambda d: d["fecha"])
    return datos, errores
