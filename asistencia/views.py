from io import BytesIO

from django.contrib import messages
from django.contrib.auth import login, logout
from django.contrib.auth.forms import AuthenticationForm
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .decorators import admin_required, estudiante_required
from .importers import parse_estudiantes, parse_fechas
from .models import Actividad, Asistencia, Estudiante, FechaClase, Modulo
from .services import (ETIQUETAS, activar_modulo, deshabilitar_modulo, eliminar_modulo, estado_calculado,
                       habilitar_modulo, marcar_asistencia, modificar_asistencia)

# ───────────────────────── Estudiante ─────────────────────────


def login_estudiante(request):
    if request.session.get("estudiante_id"):
        return redirect("mi_asistencia")
    modulo = Modulo.activo_actual()
    if request.method == "POST":
        dato = request.POST.get("identificador", "").strip()
        est = None
        if modulo and dato:
            est = (Estudiante.objects.filter(modulo=modulo)
                   .filter(Q(correo__iexact=dato) | Q(nro_registro=dato)).first())
        if est:
            request.session.cycle_key()
            request.session["estudiante_id"] = est.pk
            return redirect("mi_asistencia")
        messages.error(request, "Esos datos no corresponden a un estudiante del módulo activo. "
                                "Revisa tu correo o número de registro.")
    return render(request, "login_estudiante.html", {"modulo": modulo})


@estudiante_required
def mi_asistencia(request):
    est = request.estudiante
    hoy = timezone.localdate()
    registros = {a.fecha_clase_id: a for a in est.asistencias.all()}
    filas, presentes, faltas = [], 0, 0
    clase_hoy = None
    for f in est.modulo.fechas.all():
        reg = registros.get(f.id)
        estado = estado_calculado(f, reg, hoy)
        presentes += estado == "PRESENTE"
        faltas += estado == "FALTA"
        if f.fecha == hoy:
            clase_hoy = {"fecha": f, "registro": reg}
        filas.append({
            "fecha": f, "estado": estado, "es_hoy": f.fecha == hoy,
            "hora_marca": timezone.localtime(reg.marcado_en) if reg and reg.marcado_en else None,
            "corregido": bool(reg and reg.modificado_por_admin),
        })
    puede_marcar = bool(clase_hoy and clase_hoy["registro"] is None)
    return render(request, "mi_asistencia.html", {
        "est": est, "filas": filas, "presentes": presentes, "faltas": faltas,
        "clase_hoy": clase_hoy, "puede_marcar": puede_marcar, "hoy": hoy,
    })


@estudiante_required
@require_POST
def marcar(request):
    ok, msg = marcar_asistencia(request.estudiante)
    (messages.success if ok else messages.error)(request, msg)
    return redirect("mi_asistencia")


@require_POST
def salir(request):
    es_admin = request.user.is_authenticated
    logout(request)  # limpia también la sesión del estudiante
    return redirect("login_admin" if es_admin else "login_estudiante")


# ───────────────────────── Administrador ─────────────────────────


def login_admin(request):
    if request.user.is_authenticated and request.user.is_staff:
        return redirect("dashboard")
    form = AuthenticationForm(request, data=request.POST or None)
    if request.method == "POST":
        if form.is_valid() and form.get_user().is_staff:
            login(request, form.get_user())
            return redirect("dashboard")
        messages.error(request, "Usuario o contraseña incorrectos, o sin permisos de administrador.")
    return render(request, "panel/login.html", {"form": form})


@admin_required
def dashboard(request):
    modulo = Modulo.activo_actual()
    ctx = {"modulo": modulo}
    if modulo:
        hoy = timezone.localdate()
        total = modulo.estudiantes.count()
        clase_hoy = modulo.fechas.filter(fecha=hoy).first()
        ctx.update({
            "total_estudiantes": total,
            "total_fechas": modulo.fechas.count(),
            "clase_hoy": clase_hoy,
            "presentes_hoy": (Asistencia.objects.filter(fecha_clase=clase_hoy, estado="PRESENTE").count()
                              if clase_hoy else 0),
            "proxima": modulo.fechas.filter(fecha__gt=hoy).first(),
            "ultimas": Actividad.objects.filter(modulo=modulo)[:6],
        })
    return render(request, "panel/dashboard.html", ctx)


def _sin_modulo(request):
    messages.error(request, "No hay un módulo activo. Habilita uno existente o importa uno nuevo.")
    return redirect("modulos")


@admin_required
def estudiantes(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo(request)
    q = request.GET.get("q", "").strip()
    qs = modulo.estudiantes.all()
    if q:
        qs = qs.filter(Q(nombre__icontains=q) | Q(nro_registro__icontains=q))
    page = Paginator(qs, 25).get_page(request.GET.get("page"))
    return render(request, "panel/estudiantes.html", {"page": page, "q": q, "modulo": modulo, "qs": f"q={q}&"})


def _filas_asistencia(modulo, q, fecha_id):
    hoy = timezone.localdate()
    est_qs = modulo.estudiantes.all()
    if q:
        est_qs = est_qs.filter(Q(nombre__icontains=q) | Q(nro_registro__icontains=q))
    fechas = list(modulo.fechas.all())
    if fecha_id:
        fechas = [f for f in fechas if str(f.id) == fecha_id]
    estudiantes_l = list(est_qs)
    regs = {(a.estudiante_id, a.fecha_clase_id): a
            for a in Asistencia.objects.filter(estudiante__in=est_qs, fecha_clase__in=[f.id for f in fechas])}
    filas = []
    for f in fechas:
        for e in estudiantes_l:
            reg = regs.get((e.id, f.id))
            filas.append({
                "estudiante": e, "fecha": f, "estado": estado_calculado(f, reg, hoy),
                "hora_marca": timezone.localtime(reg.marcado_en) if reg and reg.marcado_en else None,
                "corregido": bool(reg and reg.modificado_por_admin),
            })
    return filas


@admin_required
def asistencias(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo(request)
    q = request.GET.get("q", "").strip()
    fecha_id = request.GET.get("fecha", "")
    filas = _filas_asistencia(modulo, q, fecha_id)
    totales = {
        "presentes": sum(f["estado"] == "PRESENTE" for f in filas),
        "faltas": sum(f["estado"] == "FALTA" for f in filas),
        "pendientes": sum(f["estado"] in ("PENDIENTE", "PROGRAMADA") for f in filas),
    }
    page = Paginator(filas, 25).get_page(request.GET.get("page"))
    return render(request, "panel/asistencias.html", {
        "page": page, "q": q, "fecha_id": fecha_id, "fechas": modulo.fechas.all(),
        "totales": totales, "modulo": modulo, "qs": f"q={q}&fecha={fecha_id}&",
    })


@admin_required
@require_POST
def modificar(request):
    modulo = Modulo.activo_actual()
    est = get_object_or_404(Estudiante, pk=request.POST.get("estudiante"), modulo=modulo)
    fecha = get_object_or_404(FechaClase, pk=request.POST.get("fecha"), modulo=modulo)
    ok, msg = modificar_asistencia(request.user, est, fecha, request.POST.get("estado"))
    (messages.success if ok else messages.warning)(request, msg)
    destino = request.POST.get("next", "")
    if not url_has_allowed_host_and_scheme(destino, allowed_hosts={request.get_host()}):
        destino = "asistencias"
    return redirect(destino)


@admin_required
def actividades(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo(request)
    q = request.GET.get("q", "").strip()
    fecha = request.GET.get("fecha", "").strip()
    qs = Actividad.objects.filter(modulo=modulo)
    if q:
        qs = qs.filter(Q(usuario__icontains=q) | Q(registro_afectado__icontains=q))
    if fecha:
        qs = qs.filter(fecha_hora__date=fecha)
    page = Paginator(qs, 25).get_page(request.GET.get("page"))
    return render(request, "panel/actividades.html", {
        "page": page, "q": q, "fecha": fecha, "modulo": modulo, "qs": f"q={q}&fecha={fecha}&",
    })


@admin_required
def modulo_importar(request):
    contexto = {"modulo": Modulo.activo_actual()}
    if request.method == "POST":
        nombre = request.POST.get("nombre", "").strip()
        f_est, f_fec = request.FILES.get("estudiantes"), request.FILES.get("fechas")
        errores = []
        if not nombre:
            errores.append("Indica un nombre para el nuevo módulo.")
        if not f_est:
            errores.append("Falta el Excel de estudiantes.")
        if not f_fec:
            errores.append("Falta el Excel de fechas.")
        est, fec = [], []
        if not errores:
            est, e1 = parse_estudiantes(f_est)
            fec, e2 = parse_fechas(f_fec)
            errores = [f"Estudiantes · {e}" for e in e1] + [f"Fechas · {e}" for e in e2]
        if errores:
            contexto.update({"errores": errores, "nombre": nombre})
        else:
            # RF-16: vista previa antes de activar. Se guarda en sesión hasta confirmar.
            request.session["import_pendiente"] = {"nombre": nombre, "estudiantes": est, "fechas": fec}
            return render(request, "panel/modulo_preview.html", {
                "nombre": nombre, "estudiantes": est, "fechas": fec, "modulo": contexto["modulo"],
            })
    return render(request, "panel/modulo_importar.html", contexto)


@admin_required
@require_POST
def modulo_confirmar(request):
    datos = request.session.pop("import_pendiente", None)
    if not datos:
        messages.error(request, "No hay una importación pendiente. Sube los archivos de nuevo.")
        return redirect("modulo_importar")
    modulo = activar_modulo(request.user, datos["nombre"], datos["estudiantes"], datos["fechas"])
    messages.success(request, f"Módulo «{modulo.nombre}» activado. Los datos del módulo anterior se conservan en la base de datos.")
    return redirect("dashboard")


@admin_required
def exportar(request):
    """RF-17: asistencias del módulo activo en Excel (estudiantes × fechas)."""
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo(request)
    hoy = timezone.localdate()
    fechas = list(modulo.fechas.all())
    regs = {(a.estudiante_id, a.fecha_clase_id): a
            for a in Asistencia.objects.filter(estudiante__modulo=modulo)}
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencias"
    ws.append(["Registro", "Nombre", "Correo"] + [f.fecha.strftime("%d/%m/%Y") for f in fechas]
              + ["Presentes", "Faltas"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17233B")
    for e in modulo.estudiantes.all():
        estados = [estado_calculado(f, regs.get((e.id, f.id)), hoy) for f in fechas]
        ws.append([e.nro_registro, e.nombre, e.correo] + [ETIQUETAS[s] for s in estados]
                  + [estados.count("PRESENTE"), estados.count("FALTA")])
    for col, ancho in zip("ABC", (12, 30, 30)):
        ws.column_dimensions[col].width = ancho
    buffer = BytesIO()
    wb.save(buffer)
    resp = HttpResponse(buffer.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="asistencias_{timezone.localdate():%Y%m%d}.xlsx"'
    return resp


@admin_required
def modulos(request):
    """Lista de todos los módulos, con opción de habilitar o deshabilitar."""
    lista = (Modulo.objects
             .annotate(n_estudiantes=Count("estudiantes", distinct=True),
                       n_fechas=Count("fechas", distinct=True),
                       n_asistencias=Count("estudiantes__asistencias", distinct=True))
             .order_by("-creado_en", "-id"))
    return render(request, "panel/modulos.html", {"lista": lista})


@admin_required
@require_POST
def modulo_habilitar(request, pk):
    modulo = get_object_or_404(Modulo, pk=pk)
    habilitar_modulo(request.user, modulo)
    messages.success(request, f"Módulo «{modulo.nombre}» habilitado. Solo sus estudiantes pueden ingresar ahora.")
    return redirect("modulos")


@admin_required
@require_POST
def modulo_deshabilitar(request, pk):
    modulo = get_object_or_404(Modulo, pk=pk)
    deshabilitar_modulo(request.user, modulo)
    messages.warning(request, f"Módulo «{modulo.nombre}» deshabilitado. Ningún estudiante puede ingresar hasta habilitar un módulo.")
    return redirect("modulos")


@admin_required
@require_POST
def modulo_eliminar(request, pk):
    modulo = get_object_or_404(Modulo, pk=pk)
    ok, msg = eliminar_modulo(request.user, modulo, request.POST.get("confirmacion", ""))
    (messages.success if ok else messages.error)(request, msg)
    return redirect("modulos")
